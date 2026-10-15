import datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.views.generic import ListView, TemplateView, UpdateView

from accounts.mixins import SearchablePaginatedListMixin, TeamMemberRequiredMixin

from .forms import (
    AddExistingOrderItemForm,
    OrderItemFormSet,
    PurchaseCategoryForm,
    PurchaseItemEditForm,
    QuickCreatePurchaseItemForm,
    SupplierForm,
)
from .models import OrderList, OrderListItem, PurchaseCategory, PurchaseItem, Supplier
from .selectors import build_grouped_order_rows, get_order_item_queryset, get_order_list_queryset
from . import services


def _parse_date(date_str: str) -> datetime.date:
    try:
        return datetime.date.fromisoformat(date_str)
    except ValueError:
        raise Http404("Invalid date format. Use YYYY-MM-DD.")


def _require_active_team(request):
    if not request.team or not request.membership:
        return redirect("accounts:team_select")
    return None


@login_required
def create_today_order(request):
    redirect_response = _require_active_team(request)
    if redirect_response:
        return redirect_response

    order_date = timezone.localdate()
    order, created = services.get_or_create_order(request.team, order_date, request.user)
    if created:
        if order.item_count:
            messages.success(request, "Today's order list created from the latest order.")
        else:
            messages.success(request, "Started a new blank order list for today.")
    else:
        messages.info(request, "Opened today's existing order list.")
    return redirect("ordering:order_detail", date_str=str(order.order_date))


@login_required
def create_blank_today_order(request):
    redirect_response = _require_active_team(request)
    if redirect_response:
        return redirect_response

    order_date = timezone.localdate()
    order, created = services.get_or_create_order(
        request.team,
        order_date,
        request.user,
        blank=True,
    )
    if created:
        messages.success(request, "Started a blank order list for today.")
    else:
        messages.info(request, "Opened today's existing order list.")
    return redirect("ordering:order_detail", date_str=str(order.order_date))


@login_required
def clone_order_to_today_view(request, date_str: str):
    redirect_response = _require_active_team(request)
    if redirect_response:
        return redirect_response

    source_date = _parse_date(date_str)
    source_order = get_object_or_404(OrderList, team=request.team, order_date=source_date)
    order, created = services.get_or_create_order(
        request.team,
        timezone.localdate(),
        request.user,
        source_order=source_order,
    )
    if created:
        messages.success(request, f"Today's order list created from {source_order.order_date}.")
    else:
        messages.info(request, "Opened today's existing order list.")
    return redirect("ordering:order_detail", date_str=str(order.order_date))


class OrderListView(TeamMemberRequiredMixin, ListView):
    template_name = "ordering/order_list.html"
    context_object_name = "orders"
    paginate_by = 20

    def get_queryset(self):
        return get_order_list_queryset(self.request.team)


class OrderingSetupView(TeamMemberRequiredMixin, TemplateView):
    """Manage Suppliers and Categories — the entities the supplier-first Add
    Item flow and the per-supplier Excel export are organized around.
    Delete is SET_NULL (unassigns from items) rather than blocked/cascaded,
    and there's no rename for v1 — delete + recreate covers typos."""

    template_name = "ordering/setup.html"

    def _build_context(self, supplier_form=None, category_form=None):
        return {
            "suppliers": Supplier.objects.filter(team=self.request.team),
            "categories": PurchaseCategory.objects.filter(team=self.request.team),
            "supplier_form": supplier_form or SupplierForm(team=self.request.team),
            "category_form": category_form or PurchaseCategoryForm(team=self.request.team),
        }

    def get(self, request, *args, **kwargs):
        return render(request, self.template_name, self._build_context())

    def post(self, request, *args, **kwargs):
        action = request.POST.get("action")

        if action == "add_supplier":
            form = SupplierForm(request.POST, team=request.team)
            if form.is_valid():
                supplier = form.save(commit=False)
                supplier.team = request.team
                try:
                    supplier.save()
                except IntegrityError:
                    # clean_name()'s case-insensitive pre-check can't close
                    # the window between two concurrent submits of the same
                    # name — the DB constraint is the real guard.
                    form.add_error("name", "A supplier with this name already exists.")
                    return render(request, self.template_name, self._build_context(supplier_form=form))
                messages.success(request, f"Added supplier: {supplier.name}")
                return redirect("ordering:ordering_setup")
            return render(request, self.template_name, self._build_context(supplier_form=form))

        if action == "delete_supplier":
            Supplier.objects.filter(team=request.team, id=request.POST.get("supplier_id")).delete()
            messages.success(request, "Supplier removed.")
            return redirect("ordering:ordering_setup")

        if action == "add_category":
            form = PurchaseCategoryForm(request.POST, team=request.team)
            if form.is_valid():
                category = form.save(commit=False)
                category.team = request.team
                try:
                    category.save()
                except IntegrityError:
                    form.add_error("name", "A category with this name already exists.")
                    return render(request, self.template_name, self._build_context(category_form=form))
                messages.success(request, f"Added category: {category.name}")
                return redirect("ordering:ordering_setup")
            return render(request, self.template_name, self._build_context(category_form=form))

        if action == "delete_category":
            PurchaseCategory.objects.filter(team=request.team, id=request.POST.get("category_id")).delete()
            messages.success(request, "Category removed.")
            return redirect("ordering:ordering_setup")

        return redirect("ordering:ordering_setup")


class PurchaseItemListView(TeamMemberRequiredMixin, SearchablePaginatedListMixin, ListView):
    """Only way (besides Django admin) to re-categorize or re-assign a
    supplier on an item after it's been created."""

    model = PurchaseItem
    template_name = "ordering/purchase_item_list.html"
    context_object_name = "purchase_items"
    search_fields = ["name"]

    def get_base_queryset(self):
        return (
            PurchaseItem.objects.filter(team=self.request.team)
            .select_related("category", "supplier")
            .order_by("category", "sort_order", "name")
        )


class PurchaseItemUpdateView(TeamMemberRequiredMixin, UpdateView):
    model = PurchaseItem
    template_name = "ordering/purchase_item_form.html"
    success_url = reverse_lazy("ordering:purchase_item_list")

    def get_queryset(self):
        return PurchaseItem.objects.filter(team=self.request.team)

    def get_form(self, form_class=None):
        return PurchaseItemEditForm(
            self.request.POST or None, instance=self.get_object(), team=self.request.team
        )


class OrderDetailView(TeamMemberRequiredMixin, TemplateView):
    """
    Adding items is supplier-first: Step 1 (no `?supplier=`) shows a picker
    of the team's suppliers instead of a blind, unscoped "add anything"
    form. Step 2 (`?supplier=<id>`) scopes both the existing-item dropdown
    and the quick-create form to that supplier. The param is threaded
    through every redirect/re-render in post() so adding several items in a
    row doesn't force re-picking the supplier each time.
    """

    template_name = "ordering/order_detail.html"

    def _get_order(self):
        return get_object_or_404(
            OrderList.objects.select_related("created_by", "finalized_by"),
            team=self.request.team,
            order_date=_parse_date(self.kwargs["date_str"]),
        )

    def _get_selected_supplier(self, request):
        supplier_id = request.GET.get("supplier")
        if not supplier_id:
            return None
        return Supplier.objects.filter(team=request.team, id=supplier_id).first()

    def _detail_url(self, order, selected_supplier=None):
        url = reverse("ordering:order_detail", args=[str(order.order_date)])
        if selected_supplier:
            url += f"?supplier={selected_supplier.id}"
        return url

    def _get_formset(self, order, data=None):
        return OrderItemFormSet(
            data=data,
            queryset=get_order_item_queryset(order),
        )

    def _build_context(
        self,
        *,
        order,
        formset=None,
        add_existing_form=None,
        create_item_form=None,
        selected_supplier=None,
    ):
        formset = formset or self._get_formset(order)
        if selected_supplier and add_existing_form is None:
            add_existing_form = AddExistingOrderItemForm(team=self.request.team, supplier=selected_supplier)
        if selected_supplier and create_item_form is None:
            create_item_form = QuickCreatePurchaseItemForm(team=self.request.team)

        return {
            "order": order,
            "formset": formset,
            "groups": build_grouped_order_rows(formset),
            "suppliers": Supplier.objects.filter(team=self.request.team),
            "selected_supplier": selected_supplier,
            "add_existing_form": add_existing_form,
            "create_item_form": create_item_form,
        }

    def get(self, request, *args, **kwargs):
        order = self._get_order()
        selected_supplier = self._get_selected_supplier(request)
        return render(
            request, self.template_name,
            self._build_context(order=order, selected_supplier=selected_supplier),
        )

    def post(self, request, *args, **kwargs):
        order = self._get_order()
        action = request.POST.get("action", "save")
        selected_supplier = self._get_selected_supplier(request)

        if action == "reopen":
            services.reopen_order(order)
            messages.success(request, "Order list reopened for editing.")
            return redirect(self._detail_url(order, selected_supplier))

        if action == "erase":
            try:
                services.erase_order(order)
                messages.success(request, "Order list erased.")
                return redirect("ordering:order_list")
            except (ValueError, PermissionError) as exc:
                messages.error(request, str(exc))
                return redirect(self._detail_url(order, selected_supplier))

        if action == "remove_item":
            try:
                item_id = int(request.POST.get("item_id", ""))
            except (TypeError, ValueError):
                messages.error(request, "Invalid item.")
                return redirect(self._detail_url(order, selected_supplier))
            try:
                services.remove_order_item(item_id, request.user)
                messages.success(request, "Item removed from the order list.")
            except OrderListItem.DoesNotExist:
                # Already gone — most likely a double-tap of Remove.
                messages.info(request, "Item already removed.")
            except (ValueError, PermissionError) as exc:
                messages.error(request, str(exc))
            return redirect(self._detail_url(order, selected_supplier))

        if action in {"add_existing", "add_new"} and selected_supplier is None:
            messages.error(request, "Pick a supplier first.")
            return redirect(self._detail_url(order))

        if action == "add_existing":
            add_existing_form = AddExistingOrderItemForm(
                request.POST, team=request.team, supplier=selected_supplier
            )
            if add_existing_form.is_valid():
                purchase_item = add_existing_form.cleaned_data["purchase_item"]
                try:
                    _, created = services.add_purchase_item_to_order(order, purchase_item, request.user)
                    if created:
                        messages.success(request, f"Added {purchase_item.name} to the order list.")
                    else:
                        messages.info(request, f"{purchase_item.name} is already on this order list.")
                    return redirect(self._detail_url(order, selected_supplier))
                except (ValueError, PermissionError) as exc:
                    messages.error(request, str(exc))
            create_item_form = QuickCreatePurchaseItemForm(team=request.team)
            return render(
                request,
                self.template_name,
                self._build_context(
                    order=order,
                    add_existing_form=add_existing_form,
                    create_item_form=create_item_form,
                    selected_supplier=selected_supplier,
                ),
            )

        if action == "add_new":
            create_item_form = QuickCreatePurchaseItemForm(request.POST, team=request.team)
            if create_item_form.is_valid():
                try:
                    purchase_item, _ = services.create_purchase_item_and_add_to_order(
                        order,
                        name=create_item_form.cleaned_data["name"],
                        category=create_item_form.cleaned_data["category"],
                        supplier=selected_supplier,
                        default_unit=create_item_form.cleaned_data["default_unit"],
                        user=request.user,
                    )
                    messages.success(request, f"Created and added {purchase_item.name}.")
                    return redirect(self._detail_url(order, selected_supplier))
                except IntegrityError:
                    # clean_name()'s case-insensitive pre-check can't close
                    # the window between two concurrent submits of the same
                    # name — the DB constraint is the real guard.
                    create_item_form.add_error("name", "A catalog item with this name already exists.")
                except (ValueError, PermissionError) as exc:
                    messages.error(request, str(exc))
            add_existing_form = AddExistingOrderItemForm(team=request.team, supplier=selected_supplier)
            return render(
                request,
                self.template_name,
                self._build_context(
                    order=order,
                    add_existing_form=add_existing_form,
                    create_item_form=create_item_form,
                    selected_supplier=selected_supplier,
                ),
            )

        if action in {"save", "finalize"} and order.state != OrderList.OrderState.DRAFT:
            messages.error(request, "Cannot edit a finalized order list until it is reopened.")
            return redirect(self._detail_url(order, selected_supplier))

        formset = self._get_formset(order, data=request.POST)
        if not formset.is_valid():
            messages.error(request, "Could not save the order list. Please check the highlighted fields.")
            return render(
                request,
                self.template_name,
                self._build_context(order=order, formset=formset, selected_supplier=selected_supplier),
            )

        formset.save()

        if action == "finalize":
            services.finalize_order(order, request.user)
            messages.success(request, "Order list finalized.")
        else:
            messages.success(request, "Order list draft saved.")

        return redirect(self._detail_url(order, selected_supplier))


_EXCEL_SHEET_INVALID_CHARS = str.maketrans("", "", "[]:*?/\\")


def _unique_sheet_title(raw_name: str, used_titles: set) -> str:
    """Excel sheet names: <=31 chars, no [ ] : * ? / \\, non-blank, unique
    within the workbook."""
    base = raw_name.translate(_EXCEL_SHEET_INVALID_CHARS).strip() or "Order List"
    title = base[:31]
    suffix = 2
    while title in used_titles:
        cut = 31 - len(f" ({suffix})")
        title = f"{base[:cut]} ({suffix})"
        suffix += 1
    used_titles.add(title)
    return title


@login_required
def export_order_excel(request, date_str: str):
    redirect_response = _require_active_team(request)
    if redirect_response:
        return redirect_response

    order = get_object_or_404(OrderList, team=request.team, order_date=_parse_date(date_str))

    from collections import OrderedDict
    from openpyxl import Workbook
    from openpyxl.styles import Font

    # One sheet per supplier so each can be forwarded straight to that
    # vendor, with items grouped by category inside the sheet for
    # readability. Items with no supplier set land on their own sheet
    # rather than being silently dropped or lumped in with a real supplier.
    by_supplier = OrderedDict()
    for item in get_order_item_queryset(order):
        supplier_group = by_supplier.setdefault(item.purchase_item.display_supplier, OrderedDict())
        supplier_group.setdefault(item.purchase_item.display_category, []).append(item)

    if not by_supplier:
        by_supplier["Order List"] = OrderedDict()

    wb = Workbook()
    bold = Font(bold=True)
    used_titles = set()

    for i, (supplier, categories) in enumerate(by_supplier.items()):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = _unique_sheet_title(supplier, used_titles)

        ws["A1"] = f"Order List - {order.order_date}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Supplier: {supplier}"
        ws["A3"] = f"State: {order.get_state_display()}"
        row = 5

        if not categories:
            ws.cell(row=row, column=1, value="No items on this order yet.")

        for category, items in categories.items():
            ws.cell(row=row, column=1, value=category).font = bold
            row += 1
            ws.cell(row=row, column=1, value="Item").font = bold
            ws.cell(row=row, column=2, value="Quantity").font = bold
            ws.cell(row=row, column=3, value="Note").font = bold
            row += 1

            for item in items:
                ws.cell(row=row, column=1, value=item.purchase_item.name)
                ws.cell(row=row, column=2, value=item.quantity_text)
                ws.cell(row=row, column=3, value=item.note)
                row += 1

            row += 1

        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 32

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="order_list_{order.order_date.isoformat()}.xlsx"'
    )
    wb.save(response)
    return response
